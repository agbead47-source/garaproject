# -*- coding: utf-8 -*-
"""신규 바이어 발굴 - 분석 결과 내려받기 (엑셀).

분석 화면의 숫자를 회사에 보고하려면 화면을 캡처해 붙이는 수밖에 없었다.
캡처는 **정의가 같이 안 따라간다.** "상담 전환율 23%" 한 줄만 옮겨 붙으면
그게 어느 기간인지, 고유 기업 기준인지, 실제 데이터인지 데모인지가 사라진다.
보고서에서 그게 빠지면 다음 회의에서 숫자를 다시 설명해야 한다.

그래서 이 파일은 숫자와 **집계 조건·정의를 같은 파일 안에** 담는다.

지켜야 할 것:
  - 값이 없으면 빈 칸으로 둔다. 0 으로 채우지 않는다
    (연락한 기업이 0이라 못 구한 전환율과, 정말 0%인 전환율은 다르다)
  - '계산 불가' 는 글자 그대로 적는다. 숫자 칸에 0 을 넣지 않는다
  - 실제와 데모를 한 파일에 섞지 않는다. 파일 이름에도 어느 쪽인지 적는다
  - 화면에 없는 값을 만들어 넣지 않는다. 화면과 같은 집계를 그대로 옮긴다
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import models

TITLE_FONT = Font(bold=True, size=13)
HEAD_FILL = PatternFill("solid", fgColor="F4F5F7")
HEAD_FONT = Font(bold=True, size=10)
LABEL_FONT = Font(bold=True, size=10)
NOTE_FONT = Font(size=9, color="6B7280")
WRAP = Alignment(wrap_text=True, vertical="top")

# 숫자가 아닌 값은 이렇게 적는다. 0 과 구분하려고 글자로 둔다.
NOT_AVAILABLE = "계산 불가"


def _sheet(wb, name, title, note=""):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.append([])
    if note:
        ws.append([note])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT
        ws.append([])
    return ws


def _head(ws, headers, widths):
    ws.append(headers)
    for idx in range(1, len(headers) + 1):
        cell = ws.cell(row=ws.max_row, column=idx)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _pair(ws, label, value, note=""):
    """라벨 · 값 · 비고 한 줄. 값이 None 이면 빈 칸으로 둔다."""
    ws.append([label, "" if value is None else value, note])
    ws.cell(row=ws.max_row, column=1).font = LABEL_FONT
    if note:
        ws.cell(row=ws.max_row, column=3).font = NOTE_FONT


def _pct(value):
    """비율. 못 구한 값은 0 이 아니라 '계산 불가' 다."""
    return NOT_AVAILABLE if value is None else round(value, 1)


# ---------------------------------------------------------------------------
# 시트
# ---------------------------------------------------------------------------

def _summary_sheet(wb, data, fit_mode, collected_at):
    mode_label = "실제" if data["data_mode"] == "real" else "데모"
    fit_label = dict((m["value"], m["label"]) for m in models.MODES).get(fit_mode, fit_mode)

    ws = _sheet(wb, "요약", "신규 바이어 발굴 · 분석 요약")
    _head(ws, ["집계 조건", "값", "비고"], [22, 20, 62])

    _pair(ws, "데이터 구분", mode_label + " 데이터",
          "실제와 데모는 절대 섞지 않습니다. 이 파일은 한 쪽만 담고 있습니다")
    _pair(ws, "집계 기간", "최근 {}일".format(data["funnel"]["days"]),
          "{} 이후 최초 연락한 기업만 셉니다".format(data["funnel"]["since"]))
    _pair(ws, "기준일", data["today"], "이 날짜까지의 기록으로 집계했습니다")
    _pair(ws, "내려받은 시각", collected_at, "기준일과 다를 수 있습니다")
    _pair(ws, "적합도 기준", fit_label + " 기준",
          "OEM 과 ODM 은 기준별 가중치가 다릅니다. 점수를 서로 견주지 마세요")
    _pair(ws, "시간대", data["timezone"], "")

    ws.append([])
    _head(ws, ["지표", "값", "비고"], [22, 20, 62])

    funnel = data["funnel"]
    _pair(ws, "최초 연락 기업", funnel["contacted"], "고유 기업 수입니다 (연락 건수가 아닙니다)")
    _pair(ws, "회신 기업", funnel["replied"], "무응답 {}건".format(funnel["no_reply"]))
    _pair(ws, "상담 도달 기업", funnel["meetings"], "")
    _pair(ws, "상담 전환율 (%)", _pct(funnel["rate"]),
          "상담 도달 ÷ 최초 연락. 분모가 0이면 계산 불가입니다")

    response = data["response"]
    _pair(ws, "회신 소요 중앙값 (일)",
          NOT_AVAILABLE if response["median"] is None else round(response["median"], 1),
          "제안일이 아니라 최초 연락 단계 도달일부터 셉니다")
    _pair(ws, "회신 소요 표본", response["samples"], "회신이 있는 기업만 셉니다")
    _pair(ws, "무응답", response["no_reply"], "")

    dist = data["distribution"]
    _pair(ws, "점수 산출됨", dist["scored"], "적합도를 낼 만큼 정보가 모인 후보")
    _pair(ws, "정보 부족", dist["insufficient"],
          "확인율 {:.0f}% 미만이라 확정 추천에서 뺀 후보".format(models.COVERAGE_THRESHOLD))
    _pair(ws, "점수 미산출", dist["unscored"], "확인된 기준이 없어 점수를 못 낸 후보")

    _pair(ws, "전체 후보", data["prospects"], "")
    _pair(ws, "검토 가능", data["reviewable"],
          "점수가 나왔고 확인율도 기준을 넘긴 후보")
    _pair(ws, "오늘 후속 연락", data["due_today"], "")
    _pair(ws, "기한 지난 후속 연락", data["overdue"], "")

    if data["recent"]:
        ws.append([])
        ws.append(["최근 14일에 등록된 후보 {}건은 관찰기간이 짧습니다. "
                   "전환율이 실제보다 낮게 나옵니다.".format(data["recent"])])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT


def _channel_sheet(wb, data):
    ws = _sheet(wb, "발굴 경로별", "발굴 경로별 진행",
                "모두 고유 기업 수입니다. 같은 기업에 여러 번 연락해도 한 번만 셉니다.")
    _head(ws, ["발굴 경로", "후보", "첫 연락", "회신", "상담"], [26, 10, 10, 10, 10])
    for row in data["channels"]:
        ws.append([row["channel"], row["prospects"], row["contacted"],
                   row["replied"], row["meetings"]])
    if not data["channels"]:
        ws.append(["집계할 후보가 없습니다."])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT


def _country_sheet(wb, data, fit_mode):
    fit_label = dict((m["value"], m["label"]) for m in models.MODES).get(fit_mode, fit_mode)
    ws = _sheet(wb, "국가별 적합도", "국가별 적합도 분포 ({} 기준)".format(fit_label),
                "평균은 점수가 나온 후보만 넣어 냈습니다. "
                "점수를 못 낸 후보를 0점으로 치지 않습니다.")
    _head(ws, ["소재국", "후보", "점수 산출", "평균 적합도", "평균 확인율 (%)"],
          [18, 10, 12, 14, 16])
    for row in data["distribution"]["countries"]:
        ws.append([
            row["country"], row["count"], row["scored"],
            "" if row["avg_fit"] is None else round(row["avg_fit"], 1),
            "" if row["avg_coverage"] is None else round(row["avg_coverage"], 1),
        ])
    if not data["distribution"]["countries"]:
        ws.append(["집계할 후보가 없습니다."])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT


def _definition_sheet(wb, fit_mode):
    """숫자만 옮겨 적히지 않게, 정의를 같은 파일에 둔다."""
    ws = _sheet(wb, "지표 정의", "지표 정의와 읽는 법",
                "보고서에 숫자만 옮기면 다음 회의에서 다시 설명하게 됩니다. "
                "이 시트를 같이 붙여 주세요.")
    _head(ws, ["항목", "정의 · 주의"], [22, 96])

    rows = [
        ("상담 전환율",
         "선택 기간에 최초 연락한 고유 기업 중 기준일까지 상담에 도달한 고유 기업의 비율입니다. "
         "분모가 0이면 '계산 불가' 로 적습니다. 0% 가 아닙니다."),
        ("고유 기업 기준",
         "같은 기업에 메일을 세 번 보내도 한 번만 셉니다. 연락 '건수' 가 아니라 '기업 수' 입니다."),
        ("회신 소요 중앙값",
         "최초 연락 단계에 도달한 날부터 첫 회신까지의 일수입니다. 제안서를 보낸 날이 아닙니다. "
         "회신이 없는 기업은 표본에서 빠지므로, 무응답 건수를 같이 보셔야 합니다."),
        ("적합도",
         "일치로 판정된 기준의 가중치 합 ÷ 확인된 기준의 가중치 합 × 100 입니다. "
         "수주 확률이나 구매 의향이 아닙니다. 내부 규칙에 따른 참고값입니다."),
        ("정보 확인율",
         "확인된(일치 또는 불일치) 기준의 가중치 합 ÷ 전체 가중치 × 100 입니다. "
         "{:.0f}% 미만이면 '정보 부족' 으로 보고 확정 추천에서 뺍니다.".format(
             models.COVERAGE_THRESHOLD)),
        ("점수 미산출",
         "확인된 기준이 하나도 없어 분모가 0인 경우입니다. 0점이 아니라 '점수 없음' 입니다."),
        ("OEM / ODM 기준",
         "기준별 가중치가 다릅니다. OEM 은 제품군 대응을, ODM 은 개발 요구 대응을 더 크게 봅니다. "
         "두 기준의 점수를 서로 견주지 마세요."),
        ("실제 / 데모",
         "두 데이터는 절대 섞이지 않습니다. 이 파일은 한 쪽만 담고 있습니다."),
        ("빈 칸",
         "값을 구할 수 없어 비워 둔 자리입니다. 0 으로 바꾸지 마세요."),
    ]
    for label, text in rows:
        ws.append([label, text])
        ws.cell(row=ws.max_row, column=1).font = LABEL_FONT
        ws.cell(row=ws.max_row, column=2).alignment = WRAP

    ws.append([])
    ws.append(["적합도 기준별 가중치"])
    ws.cell(row=ws.max_row, column=1).font = LABEL_FONT
    _head(ws, ["기준", "OEM", "ODM", "보는 것"], [26, 8, 8, 70])
    for row in models.CRITERIA:
        ws.append([row["label"], row["oem"], row["odm"], row["hint"]])
        ws.cell(row=ws.max_row, column=4).alignment = WRAP


# ---------------------------------------------------------------------------

def build_analytics(data, fit_mode):
    """분석 화면 그대로를 엑셀 한 벌로."""
    collected_at = models.now_seoul().strftime("%Y-%m-%d %H:%M")

    wb = Workbook()
    wb.remove(wb.active)                      # 기본 시트는 쓰지 않는다
    _summary_sheet(wb, data, fit_mode, collected_at)
    _channel_sheet(wb, data)
    _country_sheet(wb, data, fit_mode)
    _definition_sheet(wb, fit_mode)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def analytics_filename(data_mode, days):
    """파일 이름에도 어느 데이터인지 적는다. 섞여 돌아다니면 안 된다."""
    return "발굴분석_{}_{}일_{}.xlsx".format(
        "실제" if data_mode == "real" else "데모", days, models.today_iso())
