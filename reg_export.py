# -*- coding: utf-8 -*-
"""규제 대조 결과 내보내기.

영업 담당자가 이 화면을 보는 최종 목적은 **바이어에게 회신하거나
연구소에 넘기는 것**이다. 화면을 캡처해 붙이면 근거가 같이 안 따라간다.
"레티놀 0.5%는 안 됩니다" 만 옮겨지면, 연구소는 어느 조문을 보고
얼마까지 되는지를 다시 찾아야 한다.

그래서 **판정과 근거·기준일을 한 파일에** 담는다.

지켜야 할 것:
  - 판정만 적고 근거를 빼지 않는다. 조문 번호와 허용 기준을 같이 적는다
  - 실데이터와 더미를 섞지 않는다. 줄마다 어느 쪽인지 적는다
  - 빈 값은 빈 칸으로 둔다. 0 이나 '적합' 으로 채우지 않는다
  - 수동으로 쪼개 넣은 성분은 **그렇게 적는다**. 우리가 계산한 값이다
  - 시행일이 남은 규정은 따로 한 장을 떼어 둔다. 지금 적합해도 출시 때 걸린다
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TITLE_FONT = Font(bold=True, size=13)
HEAD_FILL = PatternFill("solid", fgColor="F4F5F7")
HEAD_FONT = Font(bold=True, size=10)
LABEL_FONT = Font(bold=True, size=10)
NOTE_FONT = Font(size=9, color="6B7280")
WRAP = Alignment(wrap_text=True, vertical="top")

STATUS_LABEL = {"ok": "적합", "warn": "주의", "ban": "사용 불가"}


def _sheet(wb, name, title, note=""):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.append([])
    if note:
        ws.append([note])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT
        ws.append([])
    return ws


def _head(ws, headers, widths):
    ws.append(headers)
    for idx in range(1, len(headers) + 1):
        cell = ws.cell(row=ws.max_row, column=idx)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=1)


def _pair(ws, label, value, note=""):
    ws.append([label, "" if value is None else value, note])
    ws.cell(row=ws.max_row, column=1).font = LABEL_FONT
    if note:
        ws.cell(row=ws.max_row, column=3).font = NOTE_FONT


# ---------------------------------------------------------------------------

def _summary(wb, data):
    ws = _sheet(wb, "요약", "규제 대조 결과")
    _head(ws, ["항목", "값", "비고"], [20, 34, 62])

    _pair(ws, "성분표", data["parsed"]["file_name"],
          "예시 파일입니다" if data["parsed"]["is_sample"] else "")
    _pair(ws, "고객사", data["parsed"]["customer"], "")
    _pair(ws, "제품", data["parsed"]["product"], "")
    _pair(ws, "대조 국가", data["country_label"], "")
    _pair(ws, "내보낸 시각", data["exported_at"], "이 파일을 만든 시각입니다")

    ws.append([])
    _head(ws, ["집계", "건수", "비고"], [20, 34, 62])
    _pair(ws, "대조 결과", data["counts"]["total"], "")
    _pair(ws, "주의 필요", data["counts"]["warn"], "")
    _pair(ws, "사용 불가", data["counts"]["ban"], "")
    _pair(ws, "매칭 못한 행", data["parsed"]["open_unmatched"],
          "INCI 명이 없어 판정에서 빠진 줄")
    _pair(ws, "수동으로 쪼갠 성분", data["parsed"]["split_count"],
          "복합 원료를 하위 성분으로 나눠 넣은 것")

    ws.append([])
    if data.get("eu_source"):
        src = data["eu_source"]
        _pair(ws, "EU 근거", "{} 통합본 {}".format(src["name"], src["celex"]),
              "기준일 {} · 성분 {}건 · {} 수집".format(
                  src["version_date"], src["count"], src["fetched_at"]))
    if data.get("us_source"):
        src = data["us_source"]
        _pair(ws, "미국 근거", src["name"],
              "{} · 연방규정집 기준일 {} · {} 수집".format(
                  src["parts"], src["cfr_date"], src["fetched_at"]))

    ws.append([])
    ws.append(["판정은 참고 자료입니다. 생산·납기를 걸기 전에 원문으로 확인하세요."])
    ws.cell(row=ws.max_row, column=1).font = NOTE_FONT


def _rows_sheet(wb, data):
    ws = _sheet(wb, "성분별 판정", "성분별 규제 판정",
                "허용 기준이 '해당 없음' 인 줄은 배합하지 않은 성분입니다. "
                "빈 칸은 값을 구하지 못한 자리이니 0 으로 바꾸지 마세요.")
    _head(ws, ["국가", "성분", "INCI", "표기 함량", "허용 기준", "판정",
               "근거 규정", "데이터", "출처", "메모"],
          [10, 22, 26, 16, 26, 10, 34, 10, 22, 60])

    for row in data["rows"]:
        ws.append([
            row["country_name"],
            row["name"],
            row["inci"],
            row["requested"],
            row["limit"],
            STATUS_LABEL.get(row["status"], row["status"]),
            row["rule"],
            "실데이터" if row["live"] else "더미",
            ("{} {} × {}".format(row["from_mix"], row["mix_dose"],
                                 row["mix_ratio"] or "?")
             if row.get("from_mix") else "성분표"),
            row["note"],
        ])
        ws.cell(row=ws.max_row, column=10).alignment = WRAP


def _unmatched_sheet(wb, data):
    parsed = data["parsed"]
    ws = _sheet(wb, "매칭·수동 매핑", "매칭하지 못한 행과 수동 매핑",
                "자사 혼합물·추출물 블렌드는 INCI 사전에 없어 자동으로 대조되지 "
                "않습니다. 사람이 원료사 사양서를 보고 쪼갠 값입니다.")
    _head(ws, ["원료", "원문 표기", "상태", "하위 INCI", "원료 내 비율",
               "처방 내 함량", "비고"],
          [26, 30, 12, 30, 14, 14, 50])

    for item in parsed["unmatched"]:
        if not item["resolved"]:
            ws.append([item["name"], item["raw"], "미해결", "", "", "",
                       item["note"]])
            ws.cell(row=ws.max_row, column=7).alignment = WRAP
            continue
        mix = item["mix"]
        for index, part in enumerate(mix["parts"]):
            ws.append([
                item["name"] if index == 0 else "",
                item["raw"] if index == 0 else "",
                "쪼갬" if index == 0 else "",
                part["inci"],
                part["ratio"],
                part["actual"],
                mix["ratio_note"] if index == 0 else "",
            ])
            ws.cell(row=ws.max_row, column=7).alignment = WRAP

    ws.append([])
    ws.append(["처방 내 함량 = 처방 투입량 × 원료 내 비율. "
               "비율을 적지 않은 성분은 계산하지 않고 빈 칸으로 둡니다."])
    ws.cell(row=ws.max_row, column=1).font = NOTE_FONT


def _timeline_sheet(wb, data):
    timeline = data.get("timeline") or {}
    rows = timeline.get("rows") or []
    if not rows:
        return
    ws = _sheet(wb, "시행일", "시행일이 남은 규정",
                "지금 적합해도 안심할 수 없습니다. 화장품은 개발부터 출시까지 "
                "6개월~1년이 걸려서, 지금 잡는 처방이 출시 시점에는 못 팔릴 수 "
                "있습니다.")
    _head(ws, ["지역", "규정", "내용", "상태", "신규 출시 금지", "시장 철수",
               "근거", "원문"],
          [10, 30, 46, 16, 16, 14, 34, 46])

    for row in rows:
        ws.append([row["region"], row["title"], row["what"],
                   row["state_meta"]["label"], row["new_ban"], row["pull"],
                   row["rule"], row["url"]])
        ws.cell(row=ws.max_row, column=3).alignment = WRAP

    ws.append([])
    ws.append(["이 표는 사람이 원문을 보고 수기로 정리한 것입니다. "
               "마지막 대조 {} · {}. 생산·납기를 걸기 전에 반드시 원문으로 "
               "확인하세요.".format(timeline.get("reviewed_on", ""),
                                    timeline.get("reviewed_by", ""))])
    ws.cell(row=ws.max_row, column=1).font = NOTE_FONT


def _sell_sheet(wb, data):
    board = data.get("us_board")
    if not board:
        return
    ws = _sheet(wb, "미국 판매 관점", "미국 판매 관점",
                "법에 걸리느냐와 팔 수 있느냐는 다릅니다.")
    _head(ws, ["구분", "판정", "설명", "걸리는 성분"], [22, 18, 62, 40])

    otc = board["otc"]
    ws.append(["OTC 해당 여부", otc["level_meta"]["label"], otc["why"],
               " · ".join(h["name"] for h in otc["hits"])])
    ws.cell(row=ws.max_row, column=3).alignment = WRAP

    for std in board["retail"]:
        ws.append([std["name"], std["verdict_meta"]["label"], std["why"],
                   " · ".join(h["name"] for h in std["hits"] + std["softs"])])
        ws.cell(row=ws.max_row, column=3).alignment = WRAP
        ws.append(["", "", "법이 아니라 {}가 정한 구매 조건입니다. 어겨도 불법은 "
                          "아니고 그 매장에 못 들어갑니다. 저희가 공개 문서를 보고 "
                          "정리한 배제 목록({}) 기준입니다."
                   .format(std["org"], std["as_of"]), ""])
        ws.cell(row=ws.max_row, column=3).font = NOTE_FONT
        ws.cell(row=ws.max_row, column=3).alignment = WRAP


def build(data):
    wb = Workbook()
    wb.remove(wb.active)
    _summary(wb, data)
    _rows_sheet(wb, data)
    _unmatched_sheet(wb, data)
    _sell_sheet(wb, data)
    _timeline_sheet(wb, data)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def filename(country_label, today):
    return "규제대조_{}_{}.xlsx".format(country_label.replace(" ", ""), today)
